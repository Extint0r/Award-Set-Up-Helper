import os
import re
import fitz  # PyMuPDF
import hashlib
import sqlite3
import openpyxl
import pandas as pd
from pathlib import Path
from datetime import datetime
from typing import Dict, List, Any, Optional, Tuple, Set
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font, Alignment, numbers

# Suppress PyMuPDF display errors/warnings
fitz.TOOLS.mupdf_display_errors(False)

# ==========================================
# PATH & ENVIRONMENT CONFIGURATION
# ==========================================
TRIAGE_EXCEL_PATH  = Path(r"2026_7_20_CAYUSE_ORACLE_TRIAGE.xlsx")
ALN_CSV_PATH       = Path(r"ALN.csv")

# Dual Folder Sources
OSP_SOURCE_DIR     = Path(r"D:\0-Batch-AWARDS\processed_files")
ORACLE_PARENT_DIR  = Path(r"D:\OSR pdf notices")

# Output Directories & Artifacts
MD_OUTPUT_DIR      = Path(r"D:\0-Batch-AWARDS\processed_files\Markdown")
REVIEW_DIR         = Path(r"D:\0-Batch-AWARDS\processed_files\Review")

OUTPUT_EXCEL_PATH  = REVIEW_DIR / "AUDIT_INDEX_TABLE.xlsx"
OUTPUT_SQLITE_PATH = REVIEW_DIR / "AUDIT_INDEX_TABLE.db"

# ==========================================
# REGEX PATTERNS & MATCHING ENGINES
# ==========================================
RE_CAYUSE_PROJ = re.compile(r'(?<![A-Za-z0-9])(\d{2}-\d{4})(?![A-Za-z0-9])')
RE_CAYUSE_PROP = re.compile(r'(?<![A-Za-z0-9])(A\d{2}-\d{4})(?![A-Za-z0-9])', re.IGNORECASE)
RE_ORACLE_NUM  = re.compile(r'\b([1-9]\d{5})\b')
RE_BANNER_UID  = re.compile(r'(?<![A-Za-z0-9])(R\d{4,6})(?![A-Za-z0-9])', re.IGNORECASE)
RE_ALN         = re.compile(r'(?<!\d)(\d{2}\.\d{3})(?!\d)')

RE_EXECUTION_DATE = re.compile(
    r'(?:Date|Execution\s*Date|Notice\s*Date|Award\s*Date)\s*[:\=]?\s*(\d{1,2}/\d{1,2}/\d{2,4})',
    re.IGNORECASE
)

RE_DATE_RANGE = re.compile(
    r'(?:Project\s*Period|Award\s*Period|Performance\s*Period|Grant\s*Period)\s*[:\=]?\s*'
    r'(\d{1,2}/\d{1,2}/\d{2,4})\s*[-–—\s+to\s+]+\s*(\d{1,2}/\d{1,2}/\d{2,4})',
    re.IGNORECASE
)

RE_CEILING_AMOUNT = re.compile(
    r'(?:Project\s*Total\s*Amount|Total\s*Award\s*Amount|Total\s*Ceiling|Award\s*Ceiling|Total\s*Project\s*Ceiling)\s*(?:\([^)]*\))?\s*[:\=]?\s*\$?\s*([\d,]+(?:\.\d{2})?)\s*(k|m|million|thousand)?',
    re.IGNORECASE
)

RE_OBLIGATED_ACTION = re.compile(
    r'(?:Current\s*Action.*?totaling|Obligated\s*Amount|Amount\s*Awarded\s*This\s*Action|Funding\s*This\s*Action|Action\s*Amount|This\s*Action)\s*[:\=]?\s*\$?\s*([\d,]+(?:\.\d{2})?)\s*(k|m|million|thousand)?',
    re.IGNORECASE
)

RE_BUDGET_TOTAL = re.compile(
    r'(?:Total\s*Budget|Total\s*Costs|Total\s*Amount|Total\s*Direct\s*and\s*Indirect)\s*[:\=]?\s*\$?\s*([\d,]+(?:\.\d{2})?)\s*(k|m|million|thousand)?',
    re.IGNORECASE
)

RE_INDIRECT_COST = re.compile(
    r'(?:Indirect\s*Costs?|F&A\s*Costs?|Facilities\s*&\s*Admin|Overhead)\s*[:\=]?\s*\$?\s*([\d,]+(?:\.\d{2})?)\s*(k|m|million|thousand)?',
    re.IGNORECASE
)

BUDGET_PAGE_KEYWORDS = [
    "direct cost", "f&a", "facilities and administrative", "f&a rate", 
    "f&a cost", "indirect cost", "total direct", "modified total direct"
]

# Master Crosswalk Lookups
LOOKUP_PROJ_TO_ORACLE: Dict[str, str] = {}
LOOKUP_PROP_TO_ORACLE: Dict[str, str] = {}
LOOKUP_ORACLE_TO_PROJ: Dict[str, str] = {}

# ==========================================
# HELPER FUNCTIONS & FORMATTERS
# ==========================================
def parse_dollar_amount(val_str: Any) -> float:
    """Cleans currency strings, including shorthand numbers ($500k, $1.8M)."""
    if pd.isna(val_str) or val_str is None:
        return 0.0
    s = str(val_str).strip().lower()
    
    # Shorthand matching (e.g., $500k, $1.8m)
    m_shorthand = re.search(r'\$?\s*([\d,]+(?:\.\d+)?)\s*(k|m|million|thousand|b|billion)\b', s, re.IGNORECASE)
    if m_shorthand:
        num = float(m_shorthand.group(1).replace(',', ''))
        unit = m_shorthand.group(2).lower()
        if unit in ('k', 'thousand'):
            return num * 1_000.0
        elif unit in ('m', 'million'):
            return num * 1_000_000.0
        elif unit in ('b', 'billion'):
            return num * 1_000_000_000.0
            
    try:
        clean_str = re.sub(r'[^\d.]', '', s)
        return float(clean_str) if clean_str else 0.0
    except ValueError:
        return 0.0

def parse_date(date_str: Any) -> Optional[datetime]:
    """Parses date string into standard datetime object."""
    if pd.isna(date_str) or not date_str:
        return None
    for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(date_str).strip(), fmt)
        except ValueError:
            pass
    return None

def compute_file_hash(pdf_path: Path) -> str:
    """Computes MD5 hash for exact binary duplicate matching."""
    hasher = hashlib.md5()
    try:
        with open(pdf_path, 'rb') as f:
            for chunk in iter(lambda: f.read(65536), b""):
                hasher.update(chunk)
        return hasher.hexdigest()
    except Exception:
        return ""

def discover_pdf_files() -> List[Tuple[Path, str]]:
    """Discovers PDFs across OSP folder and non-empty Oracle FY subfolders."""
    discovered = []
    
    # 1. Sweep Primary OSP Directory
    if OSP_SOURCE_DIR.exists():
        for p in OSP_SOURCE_DIR.glob("*.pdf"):
            if p.is_file():
                discovered.append((p, "OSP"))

    # 2. Sweep Oracle Subfolders (FY 2018 - FY 2026)
    if ORACLE_PARENT_DIR.exists():
        fy_subdirs = [p for p in ORACLE_PARENT_DIR.glob("FY 20*") if p.is_dir()]
        for fy_dir in sorted(fy_subdirs):
            fy_pdfs = list(fy_dir.rglob("*.pdf"))
            for p in fy_pdfs:
                if p.is_file():
                    discovered.append((p, "ORACLE"))

    return discovered

def find_markdown_file(pdf_path: Path, md_dir: Path) -> Optional[Path]:
    """Finds existing Markdown file regardless of prefix variations."""
    exact_path = md_dir / f"{pdf_path.stem}.md"
    if exact_path.exists():
        return exact_path
    matches = list(md_dir.glob(f"*{pdf_path.stem}.md"))
    if matches:
        return matches[0]
    return None

def format_master_yaml_header(metadata: dict) -> str:
    """Formats unified metadata into YAML frontmatter block."""
    lines = ["---"]
    for k, v in metadata.items():
        if isinstance(v, list):
            lines.append(f"{k}: {v}")
        elif isinstance(v, bool):
            lines.append(f"{k}: {str(v).lower()}")
        elif isinstance(v, (int, float)):
            lines.append(f"{k}: {v}")
        elif v is None:
            lines.append(f"{k}: ''")
        else:
            clean_val = str(v).replace("'", "''")
            lines.append(f"{k}: '{clean_val}'")
    lines.append("---\n\n")
    return "\n".join(lines)

def load_aln_database(aln_csv_path: Path) -> Optional[pd.DataFrame]:
    """Loads Assistance Listing Numbers (ALN) catalog if available."""
    if not aln_csv_path.exists():
        return None
    try:
        df_aln = pd.read_csv(aln_csv_path, dtype=str)
        df_aln.columns = [str(c).strip().upper() for c in df_aln.columns]
        return df_aln
    except Exception:
        return None

def load_system_baselines(triage_excel_path: Path) -> Dict[str, Dict[str, Any]]:
    """Loads baseline figures and populates lookup maps."""
    global LOOKUP_PROJ_TO_ORACLE, LOOKUP_PROP_TO_ORACLE, LOOKUP_ORACLE_TO_PROJ
    if not triage_excel_path.exists():
        return {}
    try:
        xls = pd.ExcelFile(triage_excel_path)
        sheet_to_load = 'TRIAGE' if 'TRIAGE' in xls.sheet_names else ('MASTER' if 'MASTER' in xls.sheet_names else xls.sheet_names[0])
        
        df = pd.read_excel(triage_excel_path, sheet_name=sheet_to_load)
        df.columns = [str(c).strip().upper() for c in df.columns]
        
        baselines = {}
        for _, row in df.iterrows():
            oracle_num = str(row.get("ORACLE_AWARD_NUMBER", "")).split('.')[0].strip()
            cayuse_proj = str(row.get("CAYUSE_PROJECT_NUMBER", "")).strip()
            cayuse_prop = str(row.get("CAYUSE_PROPOSAL_NUMBER", "")).strip()

            if oracle_num and oracle_num not in ('nan', 'None', ''):
                if cayuse_proj and cayuse_proj not in ('nan', 'None', ''):
                    LOOKUP_PROJ_TO_ORACLE[cayuse_proj] = oracle_num
                    LOOKUP_ORACLE_TO_PROJ[oracle_num] = cayuse_proj
                if cayuse_prop and cayuse_prop not in ('nan', 'None', ''):
                    LOOKUP_PROP_TO_ORACLE[cayuse_prop] = oracle_num

            data = {
                "ORACLE_AWARD_NUMBER": oracle_num if oracle_num not in ('nan', 'None') else "",
                "CAYUSE_PROJECT_NUMBER": cayuse_proj if cayuse_proj not in ('nan', 'None') else "",
                "CAYUSE_OBLIGATED": parse_dollar_amount(row.get("CAYUSE_OBLIGATED_TOTAL", 0)),
                "ORACLE_OBLIGATED": parse_dollar_amount(row.get("ORACLE_BUDGET_BURDENED_COST", row.get("ORACLE_HARD_LIMIT", 0))),
                "CAYUSE_TOTAL_AMOUNT": parse_dollar_amount(row.get("CAYUSE_TOTAL_AMOUNT", 0)),
                "ORACLE_HEADER_HARD_LIMIT": parse_dollar_amount(row.get("ORACLE_HEADER_HARD_LIMIT", 0))
            }

            if oracle_num and oracle_num not in ('nan', 'None', ''):
                baselines[oracle_num] = data
            if cayuse_proj and cayuse_proj not in ('nan', 'None', ''):
                baselines[cayuse_proj] = data

        return baselines
    except Exception as e:
        print(f"Error loading baseline file '{triage_excel_path}': {e}")
        return {}

# ==========================================
# PASS 1: PER-DOCUMENT EXTRACTION
# ==========================================
def process_document_pass_1(pdf_path: Path, source_tag: str, md_dir: Path, aln_df: Optional[pd.DataFrame] = None) -> Dict[str, Any]:
    """Extracts dates, financials, IDs, applies budget table fallbacks, and formats Markdown."""
    filename = pdf_path.name
    file_hash = compute_file_hash(pdf_path)
    text = ""
    md_path = find_markdown_file(pdf_path, md_dir)

    # 1. Identifier Extraction
    cay_proj_matches = RE_CAYUSE_PROJ.findall(filename)
    cay_prop_matches = RE_CAYUSE_PROP.findall(filename)
    oracle_matches   = RE_ORACLE_NUM.findall(filename)
    banner_matches   = RE_BANNER_UID.findall(filename)
    aln_matches      = RE_ALN.findall(filename)
    
    aln_number  = aln_matches[0] if aln_matches else ""
    banner_uid  = banner_matches[0].upper() if banner_matches else ""
    
    pi_match = re.match(r'^([A-Za-z]+_[A-Za-z]+)', filename)
    lead_pi  = pi_match.group(1) if pi_match else "EXTRACTED_PI"

    cay_proj   = cay_proj_matches[0] if cay_proj_matches else ""
    cay_prop   = cay_prop_matches[0] if cay_prop_matches else ""
    oracle_num = oracle_matches[0] if oracle_matches else ""

    resolved_oracle = (
        oracle_num or 
        LOOKUP_PROJ_TO_ORACLE.get(cay_proj, "") or 
        LOOKUP_PROP_TO_ORACLE.get(cay_prop, "")
    )
    
    resolved_cayuse = (
        cay_proj or 
        LOOKUP_ORACLE_TO_PROJ.get(oracle_num, "") or 
        cay_prop
    )

    cluster_key = resolved_oracle or resolved_cayuse or "UNKNOWN"

    if oracle_num:
        match_confidence = "Active Index Match"
        match_reason = f"Matched Oracle Award Number ({oracle_num})"
    elif cay_proj:
        match_confidence = "Cayuse Inference"
        match_reason = f"Matched Cayuse Project Number ({cay_proj})"
    else:
        match_confidence = "Unmatched Baseline"
        match_reason = "No direct Oracle or Cayuse ID in filename"

    # 2. Text Extraction & Page Inspection
    budget_pages: List[int] = []
    if md_path and md_path.exists():
        with open(md_path, 'r', encoding='utf-8', errors='ignore') as f:
            text = f.read()
    else:
        try:
            doc = fitz.open(pdf_path)
            pages_text = []
            for pno, page in enumerate(doc, 1):
                p_text = page.get_text()
                pages_text.append(p_text)
                if any(kw in p_text.lower() for kw in BUDGET_PAGE_KEYWORDS):
                    budget_pages.append(pno)
            text = "\n".join(pages_text)
            doc.close()
        except Exception:
            text = ""

    # Fallback Banner & ALN text search
    if not banner_uid and text:
        banner_txt = re.search(r'(?:Rice\s*Fund\s*No\.?|Rfund\s*#?)\s*[:\=]?\s*(R\d{4,6})\b', text, re.IGNORECASE)
        if banner_txt:
            banner_uid = banner_txt.group(1).upper()

    if not aln_number and text:
        cfda_txt = re.search(r'(?:CFDA|ALN)\s*Number\s*[:\=]?\s*(\d{2}\.\d{3})', text, re.IGNORECASE)
        if cfda_txt:
            aln_number = cfda_txt.group(1)

    aln_title = "N/A"
    aln_source = "N/A"
    if aln_df is not None and not aln_df.empty and aln_number:
        match = aln_df[aln_df['ALN'] == aln_number] if 'ALN' in aln_df.columns else pd.DataFrame()
        if not match.empty:
            aln_title = match.iloc[0].get('ALN_PROGRAM_TITLE', 'N/A')

    # 3. Financial & Date Extraction
    header_scope = text[:4000] if text else ""
    
    exec_match = RE_EXECUTION_DATE.search(header_scope)
    execution_date = parse_date(exec_match.group(1)) if exec_match else None
    
    range_match = RE_DATE_RANGE.search(header_scope)
    doc_start_date = parse_date(range_match.group(1)) if range_match else None
    doc_end_date = parse_date(range_match.group(2)) if range_match else None
    
    obligated_match = RE_OBLIGATED_ACTION.search(header_scope)
    delta_obligated = parse_dollar_amount(obligated_match.group(1)) if obligated_match else 0.0
    
    ceiling_match = RE_CEILING_AMOUNT.search(header_scope)
    doc_ceiling = parse_dollar_amount(ceiling_match.group(1)) if ceiling_match else 0.0

    has_budget_table = len(budget_pages) > 0 or (bool(text) and any(kw in text.lower() for kw in BUDGET_PAGE_KEYWORDS))
    table_total, table_indirect, table_direct = 0.0, 0.0, 0.0
    budget_split_status = "NO_BUDGET_TABLE"

    if has_budget_table:
        tot_match = RE_BUDGET_TOTAL.search(text)
        ind_match = RE_INDIRECT_COST.search(text)
        if tot_match:
            table_total = parse_dollar_amount(tot_match.group(1))
        if ind_match:
            table_indirect = parse_dollar_amount(ind_match.group(1))

        if "Award History Sheet" in text:
            cum_match = re.search(r'Cumulative\s*[\r\n]+\s*\$([\d,]+(?:\.\d{2})?)\s*[\r\n]+\s*\$([\d,]+(?:\.\d{2})?)\s*[\r\n]+\s*\$([\d,]+(?:\.\d{2})?)', text)
            if cum_match:
                doc_ceiling = parse_dollar_amount(cum_match.group(3))

            vertical_action_matches = re.findall(
                r'(\d{2}/\d{2}/\d{4})\s*[\r\n]+\s*(\d{2}/\d{2}/\d{4})\s*[\r\n]+\s*(\d{2}/\d{2}/\d{4})\s*[\r\n]+\s*(\d{1,2}(?:\.\d+)?%)\s*[\r\n]+\s*\$([\d,]+(?:\.\d{2})?)\s*[\r\n]+\s*\$([\d,]+(?:\.\d{2})?)\s*[\r\n]+\s*\$([\d,]+(?:\.\d{2})?)',
                text
            )
            if vertical_action_matches:
                start_str, end_str, exec_str, _, dc_str, fa_str, tot_str = vertical_action_matches[-1]
                if not doc_start_date:
                    doc_start_date = parse_date(start_str)
                if not doc_end_date:
                    doc_end_date = parse_date(end_str)
                if not execution_date:
                    execution_date = parse_date(exec_str)
                table_direct = parse_dollar_amount(dc_str)
                table_indirect = parse_dollar_amount(fa_str)
                table_total = parse_dollar_amount(tot_str)

        # Smart Fallback Engine: Override delta_obligated if budget table is present & valid
        if table_total > 0:
            if table_direct == 0.0:
                table_direct = max(0.0, table_total - table_indirect)
            
            # Fallback if header obligation was zero or a minor truncated number
            if delta_obligated == 0.0 or (delta_obligated < 1000 and table_total >= 10000) or (delta_obligated / table_total < 0.01):
                delta_obligated = table_total
                budget_split_status = "FALLBACK_TABLE_TOTAL"
            else:
                budget_split_status = "MATCHED_HEADER" if abs(table_total - delta_obligated) < 1.0 else "PARTIAL_OR_UNMATCHED"

    action_tag = "Standard Award"
    if "RPPR" in filename or "PR" in filename:
        action_tag = "PR"
    elif "Amd" in filename or "Amendment" in filename:
        action_tag = "Amd"
    elif "OtherDoc" in filename:
        action_tag = "OtherDoc"

    master_metadata = {
        "original_filename": filename,
        "file_hash": file_hash,
        "source_tag": source_tag,
        "award_cluster_key": cluster_key,
        "oracle_award_number": resolved_oracle,
        "cayuse_project_number": resolved_cayuse,
        "cayuse_proposal_number": cay_prop,
        "banner_award_uid": banner_uid,
        "lead_pi": lead_pi,
        "match_confidence": match_confidence,
        "match_reason": match_reason,
        "action_tag": action_tag,
        "aln_number": aln_number,
        "aln_program_title": aln_title,
        "aln_source": aln_source,
        "execution_date": execution_date.strftime("%Y-%m-%d") if execution_date else "",
        "doc_start_date": doc_start_date.strftime("%Y-%m-%d") if doc_start_date else "",
        "doc_end_date": doc_end_date.strftime("%Y-%m-%d") if doc_end_date else "",
        "delta_obligated": delta_obligated,
        "doc_ceiling": doc_ceiling,
        "has_budget_table": has_budget_table,
        "budget_pages": budget_pages,
        "budget_split_status": budget_split_status,
        "table_total": table_total,
        "table_direct": table_direct,
        "table_indirect": table_indirect
    }

    # 4. Generate Markdown if missing
    if not (md_path and md_path.exists()):
        md_dir.mkdir(parents=True, exist_ok=True)
        target_md_path = md_dir / f"{pdf_path.stem}.md"
        header_str = format_master_yaml_header(master_metadata)
        
        with open(target_md_path, "w", encoding="utf-8") as f:
            f.write(f"{header_str}# Document: {filename}\n\n{text}")
        md_path = target_md_path

    return {
        "PDF_Path": str(pdf_path),
        "Markdown_Path": str(md_path) if md_path else "",
        "Filename": filename,
        "FILE_HASH": file_hash,
        "SOURCE_TAG": source_tag,
        "AWARD_CLUSTER_KEY": cluster_key,
        "ORACLE_AWARD_NUMBER": resolved_oracle,
        "CAYUSE_PROJECT_NUMBER": resolved_cayuse,
        "BANNER_AWARD_UID": banner_uid,
        "EXECUTION_DATE": execution_date,
        "DOC_START_DATE": doc_start_date,
        "DOC_END_DATE": doc_end_date,
        "DELTA_OBLIGATED": delta_obligated,
        "DOC_CEILING": doc_ceiling,
        "HAS_BUDGET_TABLE": has_budget_table,
        "TABLE_TOTAL": table_total,
        "TABLE_DIRECT": table_direct,
        "TABLE_INDIRECT": table_indirect,
        "BUDGET_SPLIT_STATUS": budget_split_status
    }

# ==========================================
# DUAL-SOURCE DEDUPLICATION & MERGING ENGINE
# ==========================================
def deduplicate_and_merge_sources(extracted_docs: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Merges OSP and Oracle Post-Award document instances, assigning ORIGIN_STATUS and CROSS_REF_PATH."""
    merged_docs: List[Dict[str, Any]] = []
    
    # Tier 1: Group by Exact MD5 Binary Hash
    hash_groups: Dict[str, List[Dict[str, Any]]] = {}
    unhashed_docs = []
    
    for doc in extracted_docs:
        h = doc.get("FILE_HASH", "")
        if h:
            hash_groups.setdefault(h, []).append(doc)
        else:
            unhashed_docs.append(doc)

    processed_hashes: Set[str] = set()

    for h, group in hash_groups.items():
        processed_hashes.add(h)
        sources = {d["SOURCE_TAG"] for d in group}
        
        primary_doc = group[0].copy()
        
        if len(sources) > 1 or len(group) > 1:
            primary_doc["ORIGIN_STATUS"] = "PRESENT_IN_BOTH" if len(sources) > 1 else f"{list(sources)[0]}_DUPLICATE"
            # Cross reference path is the secondary file path
            secondary_doc = next((d for d in group if d["PDF_Path"] != primary_doc["PDF_Path"]), None)
            primary_doc["CROSS_REF_PATH"] = secondary_doc["PDF_Path"] if secondary_doc else "N/A"
        else:
            primary_doc["ORIGIN_STATUS"] = f"{primary_doc['SOURCE_TAG']}_ONLY"
            primary_doc["CROSS_REF_PATH"] = "N/A"
            
        merged_docs.append(primary_doc)

    # Tier 2: Metadata Fingerprint Alignment for remaining files
    # Fingerprint: (AWARD_CLUSTER_KEY, round(DELTA_OBLIGATED, 2), EXECUTION_DATE, DOC_START_DATE)
    fingerprints: Dict[Tuple, List[Dict[str, Any]]] = {}
    remaining_docs = unhashed_docs
    
    final_docs = []
    for doc in merged_docs:
        if doc["ORIGIN_STATUS"].endswith("_ONLY"):
            fp = (
                doc["AWARD_CLUSTER_KEY"], 
                round(doc["DELTA_OBLIGATED"], 2), 
                doc["EXECUTION_DATE"], 
                doc["DOC_START_DATE"]
            )
            # Only match if non-trivial
            if doc["AWARD_CLUSTER_KEY"] != "UNKNOWN" and doc["DELTA_OBLIGATED"] > 0:
                fingerprints.setdefault(fp, []).append(doc)
            else:
                final_docs.append(doc)
        else:
            final_docs.append(doc)

    for fp, group in fingerprints.items():
        sources = {d["SOURCE_TAG"] for d in group}
        if len(sources) > 1:
            primary_doc = group[0].copy()
            primary_doc["ORIGIN_STATUS"] = "PRESENT_IN_BOTH"
            secondary_doc = next((d for d in group if d["PDF_Path"] != primary_doc["PDF_Path"]), None)
            primary_doc["CROSS_REF_PATH"] = secondary_doc["PDF_Path"] if secondary_doc else "N/A"
            final_docs.append(primary_doc)
        else:
            final_docs.extend(group)

    return final_docs

# ==========================================
# PASS 2 & 3: SYNTHESIS & RECONCILIATION
# ==========================================
def synthesize_portfolio_pass_2(cluster_docs: List[Dict[str, Any]], system_baselines: Dict[str, Dict[str, Any]]) -> Dict[str, Any]:
    """Aggregates portfolio data and calculates financial audit verdicts."""
    if not cluster_docs:
        return {}

    cluster_key = cluster_docs[0]["AWARD_CLUSTER_KEY"]
    
    oracle_num = next((d["ORACLE_AWARD_NUMBER"] for d in cluster_docs if d["ORACLE_AWARD_NUMBER"]), "")
    cayuse_proj = next((d["CAYUSE_PROJECT_NUMBER"] for d in cluster_docs if d["CAYUSE_PROJECT_NUMBER"]), "")

    baseline = (
        system_baselines.get(oracle_num) or 
        system_baselines.get(cayuse_proj) or 
        system_baselines.get(cluster_key, {})
    )

    if not oracle_num:
        oracle_num = baseline.get("ORACLE_AWARD_NUMBER", "")
    if not cayuse_proj:
        cayuse_proj = baseline.get("CAYUSE_PROJECT_NUMBER", "")

    cayuse_ob = baseline.get("CAYUSE_OBLIGATED", 0.0)
    oracle_ob = baseline.get("ORACLE_OBLIGATED", 0.0)
    cayuse_ceiling = baseline.get("CAYUSE_TOTAL_AMOUNT", 0.0)
    oracle_ceiling = baseline.get("ORACLE_HEADER_HARD_LIMIT", 0.0)

    c_0 = max(cayuse_ceiling, oracle_ceiling)
    active_ceiling = c_0

    sorted_docs = sorted(
        cluster_docs, 
        key=lambda x: (x["EXECUTION_DATE"] or x["DOC_START_DATE"] or datetime.min)
    )

    start_dates = [d["DOC_START_DATE"] for d in sorted_docs if d["DOC_START_DATE"]]
    end_dates = [d["DOC_END_DATE"] for d in sorted_docs if d["DOC_END_DATE"]]
    
    award_start_date_truth = min(start_dates) if start_dates else None
    award_end_date_truth = max(end_dates) if end_dates else None
    
    cum_obligated = 0.0
    highest_doc_ceiling = 0.0
    
    for doc in sorted_docs:
        cum_obligated += doc["DELTA_OBLIGATED"]
        if doc["DOC_CEILING"] > 0:
            active_ceiling = doc["DOC_CEILING"]
            highest_doc_ceiling = max(highest_doc_ceiling, doc["DOC_CEILING"])

    if active_ceiling == 0.0:
        active_ceiling = c_0

    ceiling_breach = False
    audit_status = "IN_SYNC"
    
    if cum_obligated > active_ceiling and active_ceiling > 0:
        if highest_doc_ceiling >= cum_obligated:
            active_ceiling = highest_doc_ceiling
            audit_status = "CEILING_RESOLVED_BY_CLUSTER_DOC"
        else:
            ceiling_breach = True
            audit_status = "CEILING_BREACH_UNAPPROVED"

    cay_sync = (abs(cum_obligated - cayuse_ob) < 1.0)
    orc_sync = (abs(cum_obligated - oracle_ob) < 1.0)
    
    if ceiling_breach:
        reconciliation_verdict = "FLAG_CEILING_BREACH"
        discrepancy_reason = f"Cumulative PDF obligations (${cum_obligated:,.2f}) exceed active ceiling (${active_ceiling:,.2f})."
    elif cay_sync and orc_sync:
        reconciliation_verdict = "IN_SYNC"
        discrepancy_reason = "PDF Truth, Cayuse, and Oracle amounts are fully aligned."
    elif not cay_sync and orc_sync:
        reconciliation_verdict = "CAYUSE_UPDATE_REQ"
        discrepancy_reason = f"PDF Truth (${cum_obligated:,.2f}) matches Oracle, but Cayuse (${cayuse_ob:,.2f}) requires update."
    elif cay_sync and not orc_sync:
        reconciliation_verdict = "ORACLE_UPDATE_REQ"
        discrepancy_reason = f"PDF Truth (${cum_obligated:,.2f}) matches Cayuse, but Oracle (${oracle_ob:,.2f}) requires update."
    else:
        reconciliation_verdict = "BOTH_OUT_OF_SYNC"
        if cayuse_ob == 0.0 and oracle_ob == 0.0 and cum_obligated > 0:
            discrepancy_reason = f"UNMATCHED_BASELINE: Extracted PDF funds (${cum_obligated:,.2f}), but no Cayuse/Oracle baseline record was matched."
        elif cum_obligated == 0.0 and (cayuse_ob > 0 or oracle_ob > 0):
            discrepancy_reason = f"ZERO_PDF_EXTRACTION: Baseline has obligations (Cayuse: ${cayuse_ob:,.2f}, Oracle: ${oracle_ob:,.2f}), but $0 action delta extracted from PDFs."
        else:
            discrepancy_reason = f"FINANCIAL_DISCREPANCY: PDF Truth (${cum_obligated:,.2f}) differs from Cayuse (${cayuse_ob:,.2f}) and Oracle (${oracle_ob:,.2f})."

    return {
        "AWARD_CLUSTER_KEY": cluster_key,
        "ORACLE_AWARD_NUMBER": oracle_num,
        "CAYUSE_PROJECT_NUMBER": cayuse_proj,
        "DOCUMENT_COUNT": len(sorted_docs),
        "PDF_START_DATE_TRUTH": award_start_date_truth,
        "PDF_END_DATE_TRUTH": award_end_date_truth,
        "PDF_CUMULATIVE_OBLIGATED": cum_obligated,
        "PDF_ACTIVE_CEILING": active_ceiling,
        "CAYUSE_OBLIGATED": cayuse_ob,
        "ORACLE_OBLIGATED": oracle_ob,
        "CAYUSE_CEILING": cayuse_ceiling,
        "ORACLE_CEILING": oracle_ceiling,
        "CEILING_BREACH": ceiling_breach,
        "AUDIT_STATUS": audit_status,
        "RECONCILIATION_VERDICT": reconciliation_verdict,
        "DISCREPANCY_REASON": discrepancy_reason
    }

# ==========================================
# DATABASE INITIALIZATION
# ==========================================
def initialize_database(db_path: Path):
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()
    
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS tbl_Documents (
        PDF_Path TEXT PRIMARY KEY,
        Markdown_Path TEXT,
        Filename TEXT,
        ORIGIN_STATUS TEXT,
        CROSS_REF_PATH TEXT,
        AWARD_CLUSTER_KEY TEXT,
        ORACLE_AWARD_NUMBER TEXT,
        CAYUSE_PROJECT_NUMBER TEXT,
        BANNER_AWARD_UID TEXT,
        EXECUTION_DATE TEXT,
        DOC_START_DATE TEXT,
        DOC_END_DATE TEXT,
        DELTA_OBLIGATED REAL,
        DOC_CEILING REAL,
        HAS_BUDGET_TABLE INTEGER,
        TABLE_TOTAL REAL,
        TABLE_DIRECT REAL,
        TABLE_INDIRECT REAL,
        BUDGET_SPLIT_STATUS TEXT
    )
    """)
    
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS tbl_Reconciliation_Summary (
        AWARD_CLUSTER_KEY TEXT PRIMARY KEY,
        ORACLE_AWARD_NUMBER TEXT,
        CAYUSE_PROJECT_NUMBER TEXT,
        DOCUMENT_COUNT INTEGER,
        PDF_START_DATE_TRUTH TEXT,
        PDF_END_DATE_TRUTH TEXT,
        PDF_CUMULATIVE_OBLIGATED REAL,
        PDF_ACTIVE_CEILING REAL,
        CAYUSE_OBLIGATED REAL,
        ORACLE_OBLIGATED REAL,
        CAYUSE_CEILING REAL,
        ORACLE_CEILING REAL,
        CEILING_BREACH INTEGER,
        AUDIT_STATUS TEXT,
        RECONCILIATION_VERDICT TEXT,
        DISCREPANCY_REASON TEXT,
        REVIEW_STATUS TEXT DEFAULT 'PENDING',
        REVIEWER_NOTES TEXT
    )
    """)
    
    conn.commit()
    conn.close()

# ==========================================
# EXCEL FORMATTING & WORKBOOK GENERATOR
# ==========================================
def export_audit_workbook(recon_df: pd.DataFrame, doc_df: pd.DataFrame, excel_path: Path):
    """Exports audit results with Excel Tables, Clickable Hyperlinks, and Short Dates."""
    excel_path.parent.mkdir(parents=True, exist_ok=True)
    
    # Pre-process dates to standard pandas Timestamps
    for col in ["PDF_START_DATE_TRUTH", "PDF_END_DATE_TRUTH"]:
        if col in recon_df.columns:
            recon_df[col] = pd.to_datetime(recon_df[col], errors='coerce')
            
    for col in ["EXECUTION_DATE", "DOC_START_DATE", "DOC_END_DATE"]:
        if col in doc_df.columns:
            doc_df[col] = pd.to_datetime(doc_df[col], errors='coerce')

    with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
        recon_df.to_excel(writer, sheet_name="3Way_Reconciliation", index=False)
        doc_df.to_excel(writer, sheet_name="Document_Level_Detail", index=False)

    wb = openpyxl.load_workbook(excel_path)
    
    # 1. Format 3Way_Reconciliation Sheet
    ws_recon = wb["3Way_Reconciliation"]
    tab_recon = Table(displayName="Table_3Way_Reconciliation", ref=ws_recon.dimensions)
    tab_recon.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws_recon.add_table(tab_recon)
    
    header_recon = [cell.value for cell in ws_recon[1]]
    for row in range(2, ws_recon.max_row + 1):
        for col_idx, h in enumerate(header_recon, 1):
            cell = ws_recon.cell(row=row, column=col_idx)
            if "DATE" in str(h) and cell.value:
                cell.number_format = 'm/d/yyyy'

    # 2. Format Document_Level_Detail Sheet
    ws_docs = wb["Document_Level_Detail"]
    tab_docs = Table(displayName="Table_Document_Level_Detail", ref=ws_docs.dimensions)
    tab_docs.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws_docs.add_table(tab_docs)

    header_docs = [cell.value for cell in ws_docs[1]]
    path_cols = ["PDF_Path", "Markdown_Path", "CROSS_REF_PATH"]

    for row in range(2, ws_docs.max_row + 1):
        for col_idx, h in enumerate(header_docs, 1):
            cell = ws_docs.cell(row=row, column=col_idx)
            val_str = str(cell.value) if cell.value else ""
            
            # Date Formatting
            if "DATE" in str(h) and cell.value:
                cell.number_format = 'm/d/yyyy'

            # Clickable Path Hyperlink Formatting
            if str(h) in path_cols and val_str and val_str not in ("N/A", "nan", "None", ""):
                file_uri = "file:///" + val_str.replace("\\", "/")
                cell.hyperlink = file_uri
                cell.font = Font(color="0000FF", underline="single")

    wb.save(excel_path)

# ==========================================
# MAIN EXECUTION PIPELINE
# ==========================================
def main():
    print("=== STARTING DUAL-SOURCE 3-WAY RECONCILIATION PIPELINE ===")
    initialize_database(OUTPUT_SQLITE_PATH)
    
    # 1. Load System Baselines & ALN Catalog
    system_baselines = load_system_baselines(TRIAGE_EXCEL_PATH)
    aln_df = load_aln_database(ALN_CSV_PATH)
    
    # 2. Discover PDFs across OSP & Oracle FY Folders
    pdf_tuples = discover_pdf_files()
    print(f"Discovered {len(pdf_tuples)} total PDF files across sources...")
    
    if not pdf_tuples:
        print("No PDF files found across defined sources. Exiting pipeline.")
        return

    # 3. Pass 1 Ingestion
    extracted_documents = []
    for idx, (pdf_path, source_tag) in enumerate(pdf_tuples, 1):
        doc_data = process_document_pass_1(pdf_path, source_tag, MD_OUTPUT_DIR, aln_df)
        extracted_documents.append(doc_data)
        if idx % 1000 == 0 or idx == len(pdf_tuples):
            print(f" -> Processed {idx}/{len(pdf_tuples)} documents...")

    # 4. Deduplicate & Cross-Reference Folders
    merged_documents = deduplicate_and_merge_sources(extracted_documents)
    print(f"Deduplication complete: {len(merged_documents)} unique logical document actions identified.")

    # 5. Group by Award Cluster Key
    clusters: Dict[str, List[Dict[str, Any]]] = {}
    for doc in merged_documents:
        ckey = doc["AWARD_CLUSTER_KEY"]
        clusters.setdefault(ckey, []).append(doc)
        
    # 6. Pass 2 & Pass 3 Portfolio Synthesis
    reconciliation_results = []
    for ckey, doc_list in clusters.items():
        synth = synthesize_portfolio_pass_2(doc_list, system_baselines)
        reconciliation_results.append(synth)

    # 7. SQLite Staging Export
    conn = sqlite3.connect(OUTPUT_SQLITE_PATH)
    doc_df = pd.DataFrame(merged_documents)
    recon_df = pd.DataFrame(reconciliation_results)

    # Format date columns to string for SQLite
    doc_db_df = doc_df.copy()
    for col in ["EXECUTION_DATE", "DOC_START_DATE", "DOC_END_DATE"]:
        if col in doc_db_df.columns:
            doc_db_df[col] = doc_db_df[col].astype(str)
            
    recon_db_df = recon_df.copy()
    for col in ["PDF_START_DATE_TRUTH", "PDF_END_DATE_TRUTH"]:
        if col in recon_db_df.columns:
            recon_db_df[col] = recon_db_df[col].astype(str)

    # Remove temporary helper columns before database export
    if "FILE_HASH" in doc_db_df.columns:
        doc_db_df = doc_db_df.drop(columns=["FILE_HASH", "SOURCE_TAG"])

    doc_db_df.to_sql("tbl_Documents", conn, if_exists="replace", index=False)
    recon_db_df.to_sql("tbl_Reconciliation_Summary", conn, if_exists="replace", index=False)
    conn.close()

    # 8. Export Enhanced Excel Workbook
    export_audit_workbook(recon_df, doc_df, OUTPUT_EXCEL_PATH)

    print(f"\nPipeline Execution Complete!")
    print(f"-> SQLite Staged Database: {OUTPUT_SQLITE_PATH}")
    print(f"-> Enhanced Audit Workbook: {OUTPUT_EXCEL_PATH}")

if __name__ == "__main__":
    main()