import os
from pathlib import Path
import pandas as pd
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from config import (
    TRIAGE_EXCEL_PATH, OSP_SOURCE_DIR, ORACLE_PARENT_DIR, REVIEW_DIR,
    RE_ORACLE_NUM, RE_CAYUSE_PROJ, RE_BANNER_UID
)
from core.crosswalk import build_dynamic_crosswalk, clean_id_str
from core.deduplicator import stage0_binary_preflight


def build_award_file_directory():
    """
    GOAL 1: Executive Award Directory & Document Sequencing Engine
    
    - Discovers all PDF documents across OSP and Oracle FY folders.
    - Applies Stage 0 binary pre-flight deduplication.
    - Runs Banner ID extraction and cross-walk inference.
    - Groups by Award Cluster (Oracle -> Cayuse -> Banner -> UNKNOWN).
    - Outputs EXECUTIVE_AWARD_FILE_DIRECTORY.xlsx with clickable file URLs.
    """
    print("=== STARTING EXECUTIVE AWARD DIRECTORY MAPPING ===", flush=True)
    
    # 1. Discover all PDFs across OSP & Oracle FY Folders
    discovered_pdfs = []
    if OSP_SOURCE_DIR.exists():
        for p in OSP_SOURCE_DIR.glob("*.pdf"):
            discovered_pdfs.append((p, "OSP"))

    if ORACLE_PARENT_DIR.exists():
        for fy_dir in sorted(ORACLE_PARENT_DIR.glob("FY 20*")):
            for p in fy_dir.rglob("*.pdf"):
                discovered_pdfs.append((p, "ORACLE"))

    print(f"Discovered {len(discovered_pdfs):,} total PDF files across institutional sources.", flush=True)

    # 2. Stage 0 Fast Binary Pre-Flight Optimization
    unique_pdfs, binary_duplicate_map = stage0_binary_preflight(discovered_pdfs)

    # 3. Feature Extraction (Extract Cayuse, Oracle, and Banner IDs)
    extracted_docs = []
    for pdf_path, source_tag in unique_pdfs:
        filename = pdf_path.name
        
        cay_proj_matches = RE_CAYUSE_PROJ.findall(filename)
        oracle_matches   = RE_ORACLE_NUM.findall(filename)
        banner_matches   = RE_BANNER_UID.findall(filename)

        extracted_docs.append({
            "PDF_Path": str(pdf_path),
            "Filename": filename,
            "SOURCE_TAG": source_tag,
            "RAW_CAYUSE_PROJ": cay_proj_matches[0] if cay_proj_matches else "",
            "RAW_ORACLE_NUM": oracle_matches[0] if oracle_matches else "",
            "RAW_BANNER_UID": banner_matches[0].upper() if banner_matches else "",
            "FILE_HASH": ""
        })

    # 4. Dynamic Crosswalk Resolution (Infers Cayuse Project from Banner ID)
    build_dynamic_crosswalk(extracted_docs, TRIAGE_EXCEL_PATH)

    # 5. Group by Award Cluster and Sequence Documents
    clusters = {}
    for doc in extracted_docs:
        ckey = str(doc.get("AWARD_CLUSTER_KEY", "UNKNOWN"))
        clusters.setdefault(ckey, []).append(doc)

    directory_records = []

    for ckey, docs in clusters.items():
        # Sort files deterministically by mtime and filename
        docs_sorted = sorted(
            docs,
            key=lambda d: (Path(d["PDF_Path"]).stat().st_mtime if Path(d["PDF_Path"]).exists() else 0, d["Filename"])
        )

        for idx, doc in enumerate(docs_sorted, 1):
            seq_id = f"{ckey}-{idx}"
            std_filename = f"{seq_id}.pdf"
            pdf_path_str = doc.get("PDF_Path", "")
            
            file_uri = "file:///" + pdf_path_str.replace("\\", "/") if pdf_path_str else ""

            directory_records.append({
                "AWARD_CLUSTER_KEY": ckey,
                "ORACLE_AWARD_NUMBER": clean_id_str(doc.get("ORACLE_AWARD_NUMBER")),
                "CAYUSE_PROJECT_NUMBER": clean_id_str(doc.get("CAYUSE_PROJECT_NUMBER")),
                "BANNER_AWARD_NUMBER": clean_id_str(doc.get("BANNER_AWARD_NUMBER")),
                "SEQUENCE_ID": seq_id,
                "STANDARDIZED_FILENAME": std_filename,
                "ORIGINAL_FILENAME": doc.get("Filename"),
                "SOURCE_TAG": doc.get("SOURCE_TAG"),
                "MATCH_CONFIDENCE": doc.get("MATCH_CONFIDENCE"),
                "LOCAL_FILE_PATH": pdf_path_str,
                "DOCUMENT_URL": file_uri
            })

    df_dir = pd.DataFrame(directory_records)
    
    # 6. Export to Executive Excel Workbook
    REVIEW_DIR.mkdir(parents=True, exist_ok=True)
    out_excel_path = REVIEW_DIR / "EXECUTIVE_AWARD_FILE_DIRECTORY.xlsx"

    with pd.ExcelWriter(out_excel_path, engine='openpyxl') as writer:
        df_dir.to_excel(writer, sheet_name="Award_File_Directory", index=False)

    # Apply openpyxl table styling & clickable hyperlink formatting
    wb = openpyxl.load_workbook(out_excel_path)
    ws = wb["Award_File_Directory"]
    
    tab = Table(displayName="Table_Award_Directory", ref=ws.dimensions)
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws.add_table(tab)

    headers = [cell.value for cell in ws[1]]
    for row in range(2, ws.max_row + 1):
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col_idx)
            val_str = str(cell.value) if cell.value is not None else ""
            
            if str(h) == "DOCUMENT_URL" and val_str and val_str not in ("N/A", "nan", "None", ""):
                cell.hyperlink = val_str
                cell.font = Font(color="0000FF", underline="single")

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value) if cell.value is not None else ""
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

    wb.save(out_excel_path)

    print(f"\nAward File Directory successfully generated!", flush=True)
    print(f" -> Output Workbook: {out_excel_path}", flush=True)
    print(f" -> Total Mapped Files: {len(df_dir):,} across {len(clusters):,} award clusters.", flush=True)


if __name__ == "__main__":
    build_award_file_directory()