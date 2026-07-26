import openpyxl
import pandas as pd
import json
from pathlib import Path
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def sanitize_complex_types_for_excel(df: pd.DataFrame) -> pd.DataFrame:
    """
    Converts list, dict, and complex object columns (e.g., DATE_SNIPPETS, BUDGET_PAGES) 
    into clean JSON strings to prevent openpyxl export failures.
    """
    if df.empty:
        return df
    df_clean = df.copy()
    for col in df_clean.columns:
        if df_clean[col].apply(lambda x: isinstance(x, (list, dict))).any():
            df_clean[col] = df_clean[col].apply(
                lambda x: json.dumps(x, ensure_ascii=False) if isinstance(x, (list, dict)) else ("" if pd.isna(x) else str(x))
            )
    return df_clean


def export_audit_workbook(
    recon_df: pd.DataFrame, 
    header_df: pd.DataFrame, 
    tx_df: pd.DataFrame, 
    excel_path: Path
):
    """
    STAGE 5: Enhanced Output & Report Synchronization Engine
    
    Exports audit results into 3 relational sheets with professional openpyxl formatting:
    1. 3Way_Reconciliation: Executive Compliance Verdicts, Dual-Source HITL Flags & Audit Warnings
    2. Award_Headers: Macro Award Metadata & Multi-Year Ceilings
    3. Transaction_Ledger: Itemized Actions with Currency, Date, Inferred Action Highlights, HITL Flags, and Status Formatting
    """
    excel_path.parent.mkdir(parents=True, exist_ok=True)
    
    # 1. Format datetime columns cleanly
    for col in ["PDF_START_DATE_TRUTH", "PDF_END_DATE_TRUTH", "PROJECT_START_DATE", "PROJECT_END_DATE"]:
        if col in recon_df.columns:
            recon_df[col] = pd.to_datetime(recon_df[col], format='mixed', errors='coerce')
        if col in header_df.columns:
            header_df[col] = pd.to_datetime(header_df[col], format='mixed', errors='coerce')

    date_cols_tx = ["ACTION_DATE", "BUDGET_PERIOD_START", "BUDGET_PERIOD_END"]
    for col in date_cols_tx:
        if col in tx_df.columns:
            tx_df[col] = pd.to_datetime(tx_df[col], format='mixed', errors='coerce')

    # 2. Sanitize complex JSON objects
    recon_df_clean  = sanitize_complex_types_for_excel(recon_df)
    header_df_clean = sanitize_complex_types_for_excel(header_df)
    tx_df_clean     = sanitize_complex_types_for_excel(tx_df)

    # 3. Write dataframes to Excel via openpyxl engine
    with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
        recon_df_clean.to_excel(writer, sheet_name="3Way_Reconciliation", index=False)
        header_df_clean.to_excel(writer, sheet_name="Award_Headers", index=False)
        tx_df_clean.to_excel(writer, sheet_name="Transaction_Ledger", index=False)

    wb = openpyxl.load_workbook(excel_path)
    
    # Define Fills and Fonts for Status Highlighting
    fill_green   = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") # Soft Green
    font_green   = Font(color="375623", bold=True)
    
    fill_blue    = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid") # Soft Blue (Inferred Actions)
    font_blue    = Font(color="1F4E78", bold=True)

    fill_red     = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") # Soft Red/Orange
    font_red     = Font(color="C65911", bold=True)

    fill_purple  = PatternFill(start_color="EBDEFA", end_color="EBDEFA", fill_type="solid") # Soft Purple (HITL Flag)
    font_purple  = Font(color="5C2483", bold=True)
    
    fill_gray    = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid") # Soft Gray
    font_gray    = Font(color="595959", italic=True)

    fill_yellow  = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # Soft Yellow
    font_yellow  = Font(color="7F6000")

    sheet_tables = [
        ("3Way_Reconciliation", "Table_3Way_Reconciliation"),
        ("Award_Headers", "Table_Award_Headers"),
        ("Transaction_Ledger", "Table_Transaction_Ledger")
    ]

    currency_keywords = ["AMOUNT", "OBLIGATED", "CEILING", "BUDGET", "DIRECT", "INDIRECT", "TOTAL"]

    for sheet_name, table_name in sheet_tables:
        if sheet_name not in wb.sheetnames:
            continue

        ws = wb[sheet_name]
        tab = Table(displayName=table_name, ref=ws.dimensions)
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
        ws.add_table(tab)

        headers = [cell.value for cell in ws[1]]
        path_cols = ["PDF_PATH", "PDF_Path", "Markdown_Path"]

        for row in range(2, ws.max_row + 1):
            # Check if this row is an inferred action or HITL flag on Transaction_Ledger
            is_inferred_row = False
            is_hitl_row = False

            if sheet_name == "Transaction_Ledger":
                if "IS_INFERRED_ACTION" in headers:
                    inf_col_idx = headers.index("IS_INFERRED_ACTION") + 1
                    is_inferred_row = bool(ws.cell(row=row, column=inf_col_idx).value)
                if "HITL_REVIEW_REQUIRED" in headers:
                    hitl_col_idx = headers.index("HITL_REVIEW_REQUIRED") + 1
                    val = ws.cell(row=row, column=hitl_col_idx).value
                    is_hitl_row = bool(val) and str(val).lower() not in ("false", "0", "none", "")

            for col_idx, h in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col_idx)
                val_str = str(cell.value) if cell.value is not None else ""
                h_str = str(h).upper()

                # Date Formatting
                if ("DATE" in h_str or "START" in h_str or "END" in h_str) and cell.value and not isinstance(cell.value, str):
                    cell.number_format = 'm/d/yyyy'

                # Currency Formatting
                if any(kw in h_str for kw in currency_keywords) and isinstance(cell.value, (int, float)):
                    cell.number_format = '$#,##0.00'

                # Hyperlink Formatting for File Paths
                if str(h) in path_cols and val_str and val_str not in ("N/A", "nan", "None", ""):
                    file_uri = "file:///" + val_str.replace("\\", "/")
                    cell.hyperlink = file_uri
                    cell.font = Font(color="0000FF", underline="single")

                # Conditional Verdict & Status Highlights
                if h_str in ("RECONCILIATION_VERDICT", "AUDIT_STATUS"):
                    if cell.value == "IN_SYNC":
                        cell.fill = fill_green
                        cell.font = font_green
                    elif cell.value in (
                        "FLAG_CEILING_BREACH", "HUMAN_REVIEW_REQUIRED", "BOTH_OUT_OF_SYNC", 
                        "ORACLE_OUT_OF_SYNC", "CAYUSE_OUT_OF_SYNC", "DISCREPANCY_DETECTED"
                    ):
                        cell.fill = fill_red
                        cell.font = font_red
                    elif cell.value in ("FLAG_DUAL_SOURCE_CONFLICT", "HITL_HUMAN_REVIEW_REQUIRED"):
                        cell.fill = fill_purple
                        cell.font = font_purple

                if h_str == "LEDGER_ACTION_STATUS":
                    if cell.value == "PRIMARY_ACTIVE_ACTION":
                        if is_inferred_row:
                            cell.fill = fill_blue
                            cell.font = font_blue
                        else:
                            cell.fill = fill_green
                            cell.font = font_green
                    elif cell.value == "DUPLICATE_SHADOW_RECORD":
                        cell.fill = fill_gray
                        cell.font = font_gray
                    elif cell.value == "NON_BINDING_RECORD":
                        cell.fill = fill_yellow
                        cell.font = font_yellow

                # Highlight HITL Flag columns explicitly
                if h_str in ("HITL_REVIEW_REQUIRED", "DISCREPANCY_NOTE") and is_hitl_row:
                    cell.fill = fill_purple
                    cell.font = font_purple

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value) if cell.value is not None else ""
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 65)

    wb.save(excel_path)